"""PyBotchi Agents related to files."""

from asyncio import Semaphore, TaskGroup
from base64 import b64encode
from collections.abc import AsyncGenerator
from contextlib import suppress
from functools import cached_property
from io import BytesIO
from itertools import islice
from os import getenv

from cv2 import (
    COLOR_BGR2GRAY,
    IMREAD_COLOR,
    TM_CCOEFF_NORMED,
    cvtColor,
    imdecode,
    matchTemplate,
    minMaxLoc,
)
from cv2.typing import MatLike

from docx import Document
from docx.document import Document as _Document
from docx.opc.part import Part
from docx.oxml.shape import CT_Blip
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.oxml.xmlchemy import BaseOxmlElement

from fastapi import UploadFile

from langchain_core.language_models.chat_models import BaseChatModel

from numpy import frombuffer, uint8

from pybotchi import Action, ActionReturn, Context, LLM
from pybotchi.utils import apply_placeholders

from pymupdf import Page, open as pdf_open

from .utils import DocxContent


class ManipulateFilesContent(Action):
    """Manipulate Files Content."""

    ####################################################################################################
    #                                             EXECUTION                                            #
    ####################################################################################################

    async def pre(self, context: Context) -> ActionReturn:
        """Executre pre process."""
        chat = context.llm
        if self.__temperature__ is not None:
            chat = chat.with_config(
                configurable={"llm_temperature": self.__temperature__}
            )

        files = "\n---\n".join(
            f"```{key}\n{val}\n```"
            for key, val in (await self.extract(context)).items()
        )

        response = await chat.ainvoke(
            [
                {
                    "content": apply_placeholders(self._operation_prompt, files=files),
                    "role": "system",
                },
                *islice(context.prompts, 1, None),
            ]
        )
        await context.add_response(self, response.content)

        return ActionReturn.GO

    ####################################################################################################
    #                                             UTILITIES                                            #
    ####################################################################################################

    @cached_property
    def _max_running_image_captioning(self) -> int:
        """Retrieve max running image captioning value.

        You may override this to meet your requirements.
        By default, it gets from environment variable with default value 3.
        """
        return int(getenv("MAX_RUNNING_IMAGE_CAPTIONING", "3"))

    @cached_property
    def _image_to_text_prompt(self) -> str:
        """Retrieve image to text prompt.

        You may override this to meet your requirements.
        """
        return """
You are an expert image analyst. Your task is to provide detailed, comprehensive descriptions of images.

## Instructions:
- Describe everything visible in the image systematically
- Start with the overall scene/setting, then move to specific details
- Include: objects, people, animals, colors, lighting, composition, mood/atmosphere
- Note text, signs, or writing if present
- Describe spatial relationships (foreground, background, left, right)
- Mention artistic style, photography technique, or medium if relevant
- Be objective and factual - avoid assumptions about context outside the image
- Use clear, descriptive language accessible to someone who cannot see the image

## Structure your response:
1. **Image Type**: Identify if it's a photograph, chart, diagram, document, etc.
2. **Main Content**: Key subjects, data, or focal points
3. **Text/Data**: Transcribe all visible text, numbers, labels, titles
4. **Visual Details**: Colors, layout, formatting, symbols
5. **Data Analysis** (for charts/graphs): Trends, patterns, key insights from the data
6. **Context/Background**: Supporting elements and overall composition

Provide enough detail that someone could visualize the image based on your description alone.
""".strip()

    @cached_property
    def _operation_prompt(self) -> str:
        """Retrieve operation prompt.

        You may override this to meet your requirements.
        """
        return """
You are a helpful and intelligent assistant that specializes in managing and manipulating the contents of files.
You can read, interpret, and transform text-based file contents into different formats or styles to meet the user's needs.
Your capabilities include (but are not limited to):
- Summarizing large amounts of text clearly and concisely.
- Extracting key points, data, or structured information.
- Creating tables, charts, or lists from unstructured content.
- Rewriting or reformatting content to match specific styles.
- Merging, splitting, or reorganizing sections of text.
- Converting between formats (e.g., plain text → Markdown tables).
- Providing suggestions for improving clarity and structure.
When responding:
- Always confirm your understanding of the user's request before making major transformations.
- Preserve important details and context unless the user specifies otherwise.
- Ask clarifying questions if the input is ambiguous.
- Maintain accuracy, and avoid fabricating content.
Your goal is to help the user efficiently work with and transform file contents into useful, well-structured outputs.

Files:
${files}
""".strip()

    async def get_files(
        self, context: Context
    ) -> AsyncGenerator[tuple[int, str, bytes, str], None]:
        """Retrieve files from current context.

        You may override this to meet your requirements.
        By default, it assume the file from context is from FastAPI Endpoint.
        """
        # from metadata or override context to include `files` field
        files: list[UploadFile] = context.metadata["uploads"]
        return (
            (
                id,
                file.filename or f"file-{id}",
                await file.read(),
                file.content_type or "application/octet-stream",
            )
            for id, file in enumerate(files)
        )

    async def get_ignored_images(self, context: Context) -> list[MatLike]:
        """Retrieve ignored images from current context.

        You may override this to meet your requirements.
        """
        images: list[UploadFile] = context.metadata["ignored_images"]
        return [
            cvtColor(img, COLOR_BGR2GRAY) if len(img.shape) == 3 else img
            for image in images
            if (img := imdecode(frombuffer(await image.read(), uint8), IMREAD_COLOR))
        ]

    async def extract(self, context: Context) -> dict[str, str]:
        """Extract content from file.

        You may override this to meet your requirements.
        You may also use `await context.run_in_thread` to run extract_pdf/docs in new thread
        To avoid blocking main thread when content gets big.
        """
        extracted: dict[str, str] = {}
        async for id, file_name, content, content_type in await self.get_files(context):
            match content_type:
                case "application/pdf":
                    extracted[f"{id}-{file_name}"] = await self.extract_pdf(
                        context, id, content
                    )
                case "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                    extracted[f"{id}-{file_name}"] = await self.extract_docx(
                        context, id, content
                    )
                case _:
                    pass
        return extracted

    async def extract_pdf(self, context: Context, id: int, content: bytes) -> str:
        """Extract content from pdf file."""
        text = ""
        images: dict[str, bytes] = {}
        page: Page

        ignored_images = await self.get_ignored_images(context)

        with pdf_open("pdf", content) as pdf:
            for page_num, page in enumerate(pdf.pages()):
                blocks = page.get_text("dict")["blocks"]
                for order, block in enumerate(blocks):
                    if block["type"] == 1:
                        if block["image"] and (
                            not ignored_images
                            or not await self.skip_image(
                                context, ignored_images, block["image"]
                            )
                        ):
                            placeholder = (
                                f'<img src="{id}-{page_num}-{order}.{block["ext"]}">'
                            )
                            images[placeholder] = block["image"]
                            text += f"\n{placeholder}\n\n"
                    else:
                        for line in block["lines"]:
                            text += f'{" ".join(
                                span["text"] for span in line["spans"]
                            ).strip()}\n'

        return await self.append_images_context(context, text, images)

    async def extract_docx(self, context: Context, id: int, content: bytes) -> str:
        """Extract content from docx file."""
        await self.traversing_docx_extractor(
            context,
            document := Document(BytesIO(content)),
            document.element,
            docx_content := DocxContent(
                id=id, ignored_images=await self.get_ignored_images(context)
            ),
        )

        return await self.append_images_context(
            context, docx_content.text, docx_content.images
        )

    async def traversing_docx_extractor(
        self,
        context: Context,
        document: _Document,
        element: BaseOxmlElement,
        docx_content: DocxContent,
    ) -> None:
        """Extract text from element."""
        match element:
            case CT_P():
                if text := element.text.strip():
                    docx_content.append(text if docx_content.table else f"{text}\n")
                children = element.getchildren()
                if children:
                    for child in children:
                        await self.traversing_docx_extractor(
                            context, document, child, docx_content
                        )
            case CT_Tbl():
                table_content: list[list[str]] = []
                for row in element.tr_lst:
                    row_content: list[str] = []
                    for cell in row.tc_lst:
                        cell_content = docx_content.child()
                        for child in cell.getchildren():
                            await self.traversing_docx_extractor(
                                context, document, child, cell_content
                            )
                        row_content.append(cell_content.text)
                        docx_content.merge_images(cell_content.images)
                    table_content.append(row_content)

                border = f'| {" --- |" * len(table_content[0])}'
                docx_content.append(f"\n{border}\n")
                for row_content in table_content:
                    docx_content.append(f'| {" | ".join(row_content)} |\n')
                docx_content.append("\n")
            case CT_Blip():
                if (
                    (part := document.part.rels.get(element.embed))
                    and isinstance(target := part._target, Part)
                    and (
                        not docx_content.ignored_images
                        or not self.skip_image(
                            context, docx_content.ignored_images, target.blob
                        )
                    )
                ):
                    placeholder = f'<img src="{docx_content.id}-{element.embed}-{target.filename}">'
                    docx_content.append(f"\n{placeholder}\n")
                    docx_content.add_image(placeholder, target.blob)
            case _:
                for child in element.getchildren():
                    await self.traversing_docx_extractor(
                        context, document, child, docx_content
                    )

    async def append_images_context(
        self, context: Context, content: str, images: dict[str, bytes]
    ) -> str:
        """Append images context."""
        semaphore = Semaphore(self._max_running_image_captioning)
        async with TaskGroup() as tg:
            tasks = [
                tg.create_task(
                    self.get_image_context(
                        context,
                        semaphore,
                        source,
                        image,
                    )
                )
                for source, image in images.items()
            ]

        for task in tasks:
            src, caption = task.result()
            content = content.replace(
                src,
                f"<figure>\n\t{src}\n\t<figcaption>\n{caption}\n\t</figcaption>\n</figure>",
            )

        return content.strip()

    async def get_image_context(
        self,
        context: Context,
        sem: Semaphore,
        name: str,
        content: bytes,
    ) -> tuple[str, str]:
        """Generate image context."""
        async with sem:
            response = await LLM.get("vision", BaseChatModel).ainvoke(
                [
                    {"role": "system", "content": self._image_to_text_prompt},
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "image_url",
                                "image_url": {
                                    "url": f'data:image/jpeg;base64,{b64encode(content).decode("utf-8")}',
                                    "detail": "high",
                                },
                            }
                        ],
                    },
                ]
            )
            return name, response.content

    async def skip_image(
        self,
        content: Context,
        excluded: list[MatLike],
        image: bytes,
        threshold: float = 0.8,
    ) -> bool:
        """Check if image should be skipped."""
        with suppress(Exception):
            if (c_img := imdecode(frombuffer(image, uint8), IMREAD_COLOR)) is None:
                return False

            if len(c_img.shape) == 3:
                c_img = cvtColor(c_img, COLOR_BGR2GRAY)

            return any(
                True
                for x_img in excluded
                if (
                    x_img.shape[0] <= c_img.shape[0]
                    and x_img.shape[1] <= c_img.shape[1]
                )
                and minMaxLoc(matchTemplate(c_img, x_img, TM_CCOEFF_NORMED))[1]
                >= threshold
            )
        return False
